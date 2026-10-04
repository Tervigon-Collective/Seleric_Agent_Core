#!/usr/bin/env python3
"""Metric inventory for the LIVE semantic v2 surface -> doc/METRIC_INVENTORY.xlsx.

  SELERIC_CUBE_ENV=/opt/seleric/mage-ai/infra/cube/.env \\
    uv run --with openpyxl python scripts/build_metric_inventory_v2.py [--agent /opt/seleric/Seleric_Agent]
        [--out doc/METRIC_INVENTORY.xlsx] [--brand 20] [--no-live]

Everything is read from what serves the agent today, never restated by hand:
  catalogue_v2/ (what the MCP resolves), Cube v2 :4002 (/meta, generated SQL, live values),
  ClickHouse system.tables (serve -> gold lineage), Seleric_Agent config/metric_registry.yaml,
  scripts/v2_gates.py (resolution + gates) and scripts/v2_parity.py (v1 -> v2 values, v1 Cube :4001).

Sheets: README · Metrics · Metric Variants · Retired v1 ids · Views · Dimensions · Hierarchies · Concepts ·
Glossary · Agent Registry · Resolution · Lineage · Gates.
"""
from __future__ import annotations

import argparse
import base64
import concurrent.futures as cf
import datetime
import json
import os
import re
import subprocess
import sys
import tempfile
import urllib.parse
import urllib.request
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "scripts"))

from v2_gates import compute, filter_members, filter_pairs, load_surface  # noqa: E402
from v2_parity import CUBE_ENV, V2_URL, last_full_months, load, secret, token  # noqa: E402

REF = re.compile(r"\b(gold|serve|semantic|mart|dw|atomic)\.`?(\w+)`?")


# ---------------------------------------------------------------- sources

def cube_sql(tok: str, member: str, axis: str, brand: str, start: str, end: str) -> str:
    view = member.split(".")[0]
    q = {"measures": [member], "timezone": "Asia/Kolkata",
         "filters": [{"member": f"{view}.brand_id", "operator": "equals", "values": [brand]}],
         "timeDimensions": [{"dimension": axis, "dateRange": [start, end]}]}
    req = urllib.request.Request(f"{V2_URL}/cubejs-api/v1/sql?query=" + urllib.parse.quote(json.dumps(q)),
                                 headers={"Authorization": tok})
    try:
        d = json.load(urllib.request.urlopen(req, timeout=120))
        return (d.get("sql") or {}).get("sql", [""])[0]
    except Exception as e:  # noqa: BLE001
        return f"ERR {e}"


def clickhouse_objects() -> dict[str, dict]:
    """serve / semantic / gold objects -> engine + direct upstream refs (from their DDL)."""
    env = {}
    for line in open(CUBE_ENV):
        if "=" in line and not line.lstrip().startswith("#"):
            k, v = line.strip().split("=", 1)
            env[k] = v.strip("'\"")
    q = ("SELECT database, name, engine, create_table_query FROM system.tables "
         "WHERE database IN ('serve','semantic','gold') FORMAT JSONEachRow")
    auth = base64.b64encode(f"{env['CUBEJS_DB_USER']}:{env['CUBEJS_DB_PASS']}".encode()).decode()
    req = urllib.request.Request(f"http://{env['CUBEJS_DB_HOST']}:{env.get('CUBEJS_DB_PORT', '8123')}/",
                                 data=q.encode(), headers={"Authorization": "Basic " + auth})
    out = {}
    for line in urllib.request.urlopen(req, timeout=120).read().decode().splitlines():
        r = json.loads(line)
        fq = f"{r['database']}.{r['name']}"
        to = re.search(r"\bTO\s+`?(\w+)`?\.`?(\w+)`?", r["create_table_query"]) if r["engine"] == "MaterializedView" else None
        target = f"{to.group(1)}.{to.group(2)}" if to else None
        refs = sorted({f"{a}.{b}" for a, b in REF.findall(r["create_table_query"])} - {fq, target})
        out[fq] = {"engine": r["engine"], "refs": refs, "writes": target}
    # a refreshable MV writes its target table: the table's upstream is the MV
    for fq, o in list(out.items()):
        if o.get("writes") in out:
            out[o["writes"]]["refs"] = sorted(set(out[o["writes"]]["refs"]) | {fq})
    return out


def gold_sources(obj: str, objs: dict, seen: set | None = None) -> set[str]:
    seen = seen if seen is not None else set()
    if obj in seen:
        return set()
    seen.add(obj)
    if obj.startswith("gold.") or obj.startswith("mart."):
        return {obj}
    found = set()
    for r in objs.get(obj, {}).get("refs", []):
        found |= gold_sources(r, objs, seen)
    return found


def parity(brand: str) -> dict[str, dict]:
    with tempfile.NamedTemporaryFile(suffix=".json") as f:
        subprocess.run([sys.executable, str(ROOT / "scripts" / "v2_parity.py"), "--brands", brand,
                        "--months", "1", "--json", f.name], check=False, capture_output=True)
        try:
            cells = json.load(open(f.name))
        except Exception:  # noqa: BLE001
            return {}
    return {c["old"]: c for c in cells}


# ---------------------------------------------------------------- workbook

def write_sheet(wb, title: str, header: list[str], rows: list[list], widths: dict[int, int] | None = None,
                note: str | None = None) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(italic=True, color="555555")
        r0 = 2
    for i, h in enumerate(header, 1):
        c = ws.cell(r0, i, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(vertical="top", wrap_text=True)
    for row in rows:
        ws.append([cell if not isinstance(cell, (list, dict, set, tuple)) else _flat(cell) for cell in row])
    ws.freeze_panes = ws.cell(r0 + 1, 2)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(header))}{r0 + len(rows)}"
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, min(max(len(h) + 2, 12), 40))


def _flat(v) -> str:
    if isinstance(v, dict):
        return ", ".join(f"{k}={x}" for k, x in v.items())
    return ", ".join(str(x) for x in v)


def num(v):
    return v if isinstance(v, (int, float)) else (None if v is None else str(v))


# ---------------------------------------------------------------- main

def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--agent", default=os.environ.get("SELERIC_AGENT", "/opt/seleric/Seleric_Agent"))
    ap.add_argument("--out", default=str(ROOT / "doc" / "METRIC_INVENTORY.xlsx"))
    ap.add_argument("--brand", default="20")
    ap.add_argument("--no-live", action="store_true", help="skip Cube / ClickHouse (catalogue-only sheets)")
    args = ap.parse_args()
    from openpyxl import Workbook

    live = not args.no_live
    svc, reg = load_surface(args.agent)
    cat = svc.cat
    months = last_full_months(2)
    (p_s, p_e), (s, e) = months
    tok = token(secret()) if live else ""

    def axis_of(m) -> str:
        return m.cube_mapping.time_dimension or f"{m.cube_mapping.view}.{cat.views[m.cube_mapping.view].date_dimension}"

    # ---- live values + generated SQL per metric
    values: dict[str, tuple] = {}
    sqls: dict[str, str] = {}
    if live:
        def one(m):
            if not m.is_queryable:
                return m.id, (None, None), ""
            cur = load(V2_URL, tok, m.cube_mapping.measure, axis_of(m), args.brand, s, e)
            prev = load(V2_URL, tok, m.cube_mapping.measure, axis_of(m), args.brand, p_s, p_e)
            return m.id, (cur, prev), cube_sql(tok, m.cube_mapping.measure, axis_of(m), args.brand, s, e)
        with cf.ThreadPoolExecutor(6) as ex:
            for mid, v, q in ex.map(one, cat.metrics.values()):
                values[mid], sqls[mid] = v, q
    objs = clickhouse_objects() if live else {}
    serve_of = {mid: sorted(set(re.findall(r"serve\.`?(\w+)`?", q))) for mid, q in sqls.items()}

    # ---- cross references
    formerly: dict[str, list[str]] = {}
    for d in cat.deprecations:
        if d.new:
            formerly.setdefault(d.new, []).append(d.old + (f" [{_flat(d.filters)}]" if d.filters else ""))
    gloss_by: dict[str, list[str]] = {}
    for t in cat.glossary:
        if t.canonical_id:
            gloss_by.setdefault(t.canonical_id, []).append(t.term + (f" [{_flat(t.filter)}]" if t.filter else ""))
    concept_by: dict[str, list[str]] = {}
    for c in cat.concepts.values():
        for r in c.resolves:
            if getattr(r, "metric", None):
                when = getattr(r, "when", None) or {}
                concept_by.setdefault(r.metric, []).append(c.id + (f" ({_flat(when)})" if when else ""))
    reg_by: dict[str, list[str]] = {}
    for r in reg:
        if r.get("catalogue_metric"):
            reg_by.setdefault(r["catalogue_metric"], []).append(
                r["id"] + (f" [{_flat(r['catalogue_filters'])}]" if r.get("catalogue_filters") else ""))
    hier_by_dim = {lvl: hid for hid, h in cat.hierarchies.items() for lvl in h.levels}
    modules_by_view: dict[str, list[str]] = {}
    mods = yaml.safe_load((ROOT / "catalogue_v2" / "modules.yaml").read_text())["modules"]
    for mid_, mod in mods.items():
        for v in mod.get("extra_views") or []:
            modules_by_view.setdefault(v, []).append(mid_)

    wb = Workbook()
    wb.remove(wb.active)

    # ---- README (filled last, but first sheet)
    readme = wb.create_sheet("README")

    # ---- Metrics
    rows = []
    for m in sorted(cat.metrics.values(), key=lambda x: (x.cube_mapping.view, x.id)):
        cur, prev = values.get(m.id, (None, None))
        gold = sorted(set().union(*[gold_sources(f"serve.{o}", objs) for o in serve_of.get(m.id, [])])) if objs else []
        rows.append([
            m.id, m.display_name, m.status, m.version, m.category, modules_by_view.get(m.cube_mapping.view, []),
            m.cube_mapping.view, m.cube_mapping.measure, axis_of(m), m.grain, m.aggregation, m.unit,
            m.formula.human_readable, m.formula.depends_on, m.description,
            [f"{b.name}→{b.view}" + (f" (slice {b.slice_dimension})" if b.slice_dimension else "") for b in m.bindings],
            {k: "|".join(v) for k, v in m.valid_for.items()}, m.unavailable_reason or "",
            serve_of.get(m.id, []), gold, formerly.get(m.id, []), concept_by.get(m.id, []),
            gloss_by.get(m.id, []), reg_by.get(m.id, []), m.data_owner,
            num(cur), num(prev), sqls.get(m.id, ""),
        ])
    write_sheet(wb, "Metrics", [
        "metric_id", "display_name", "status", "version", "category", "modules", "view", "cube_member", "date_axis",
        "grain", "aggregation", "unit", "formula", "depends_on", "description", "bindings", "valid_for",
        "unavailable_reason", "serve_objects", "gold_sources", "formerly_v1_ids", "concepts", "glossary_terms",
        "agent_registry_entries", "data_owner", f"value_{s[:7]}_brand{args.brand}", f"value_{p_s[:7]}_brand{args.brand}",
        "cube_sql",
    ], rows, {1: 30, 2: 28, 8: 38, 13: 50, 15: 60, 28: 80},
        note=f"One row per v2 metric id (one id = one number). Values: brand {args.brand}, IST months; ratios recomputed.")

    # ---- Metric variants (metric + filters the surface can emit)
    pairs = filter_pairs(svc, reg)
    vrows = []
    for (mid, flt), srcs in sorted(pairs.items()):
        m = cat.metrics.get(mid)
        val = None
        if live and m is not None and m.is_queryable:
            members, err = filter_members(svc, m, flt)
            val = err or load(V2_URL, tok, m.cube_mapping.measure, axis_of(m), args.brand, s, e, members)
        vrows.append([mid, dict(flt), len(srcs), sorted(srcs), num(val)])
    write_sheet(wb, "Metric Variants", ["metric_id", "filters", "n_sources", "sources", f"value_{s[:7]}"], vrows,
                {1: 28, 2: 40, 4: 90},
                note="Every metric + filter combination the catalogue (glossary, concept axes, retired-id replacements) "
                     "or the agent registry can send. Scope / platform / paid are filters, never separate ids.")

    # ---- Retired v1 ids
    par = parity(args.brand) if live else {}
    rrows = []
    for d in sorted(cat.deprecations, key=lambda x: x.old):
        if d.old in cat.metrics:
            continue  # kept id (redefined in place), not retired
        p = par.get(d.old, {})
        if d.new not in cat.metrics:
            status = "n/a (no replacement)"
        elif not p:
            status = ""
        elif isinstance(p.get("v1"), str):
            status = "not comparable (v1 query has no date axis / slice)"
        else:
            status = "equal" if p.get("match") else ("known diff" if p.get("known_diff") else "DIFF")
        rrows.append([d.old, d.new if d.new in cat.metrics else "(none — retired without replacement)", d.filters,
                      d.reason, num(p.get("v1")), num(p.get("v2")), status, p.get("known_diff") or ""])
    write_sheet(wb, "Retired v1 ids", ["v1_id", "v2_replacement", "filters", "reason", f"v1_value_{s[:7]}",
                                       f"v2_value_{s[:7]}", "parity", "known_diff"], rrows, {1: 36, 2: 30, 4: 60, 8: 50},
                note="Hard cut: the MCP rejects every v1 id with an error naming the replacement + filters. "
                     "Values: v1 Cube (:4001) vs v2 Cube (:4002), scripts/v2_parity.py.")

    # ---- Views
    by_view: dict[str, list] = {}
    for m in cat.metrics.values():
        by_view.setdefault(m.cube_mapping.view, []).append(m.id)
    view_rows = []
    for v in cat.views.values():
        serve = sorted({o for mid in by_view.get(v.name, []) for o in serve_of.get(mid, [])})
        dims = [d.id for d in cat.dimensions.values() if v.name in d.views]
        view_rows.append([v.name, v.title, modules_by_view.get(v.name, []), v.date_dimension, v.datetime_dimension,
                          len(by_view.get(v.name, [])), len(dims), serve, v.freshness.expected_cadence, v.description])
    write_sheet(wb, "Views", ["view", "title", "modules", "date_dimension", "datetime_dimension", "metrics",
                              "dimensions", "serve_objects", "cadence", "description"], view_rows, {10: 90})

    # ---- Dimensions
    write_sheet(wb, "Dimensions", ["dimension_id", "display_name", "hierarchy", "views", "aliases", "allowed_values",
                                   "unsupported_values", "description"],
                [[d.id, d.display_name, hier_by_dim.get(d.id, ""), sorted(d.views), d.aliases, d.allowed_values or [],
                  d.unsupported_values, d.description] for d in sorted(cat.dimensions.values(), key=lambda x: x.id)],
                {4: 60, 8: 70})

    # ---- Hierarchies
    write_sheet(wb, "Hierarchies", ["hierarchy", "display_name", "levels (coarse → fine)", "views", "note"],
                [[h.id, h.display_name, " → ".join(h.levels), h.views, h.note or ""] for h in cat.hierarchies.values()],
                {3: 60, 5: 60})

    # ---- Concepts
    crow = []
    for c in sorted(cat.concepts.values(), key=lambda x: x.id):
        res = [(_flat(getattr(r, "when", None) or {}) or "default") + " → " + r.metric for r in c.resolves
               if getattr(r, "metric", None)]
        axes = {k: "|".join(a.values) + (f" (default {a.default})" if a.default else "") for k, a in c.axes.items()}
        crow.append([c.id, c.display_name, c.aliases, axes, res,
                     {ax: {v: _flat(f) for v, f in vals.items()} for ax, vals in c.axis_filters.items()}])
    write_sheet(wb, "Concepts", ["concept", "display_name", "aliases", "axes", "resolves", "axis_filters"], crow,
                {3: 50, 4: 60, 5: 70, 6: 70})

    # ---- Glossary
    write_sheet(wb, "Glossary", ["term", "metric_id", "filter", "dimension_id", "definition"],
                [[t.term, t.canonical_id or "", t.filter, t.canonical_dimension_id or "", t.definition or ""]
                 for t in sorted(cat.glossary, key=lambda x: x.term.lower())], {1: 34, 2: 28, 5: 70})

    # ---- Agent registry
    write_sheet(wb, "Agent Registry", ["registry_id", "catalogue_metric", "catalogue_filters", "aliases", "domain",
                                       "owner", "in_catalogue_v2"],
                [[r["id"], r.get("catalogue_metric") or "", r.get("catalogue_filters") or {}, r.get("aliases") or [],
                  r.get("domain") or "", r.get("owner") or "", bool(r.get("catalogue_metric") in cat.metrics)]
                 for r in reg], {1: 40, 2: 28, 3: 40, 4: 50},
                note=f"Seleric_Agent config/metric_registry.yaml ({args.agent}).")

    # ---- Resolution + gates
    g = compute(svc, reg, reg, live)
    write_sheet(wb, "Resolution", ["phrase", "resolve_term", "resolve_concept", "search_top1", "agent_registry",
                                   "verdict"],
                [[r["phrase"], r["resolve_term"], r["resolve_concept"], r["search_top1"], r["agent_registry"],
                  r["verdict"]] for r in g["rows"]], {1: 36},
                note="Every phrase through each resolver; CONFLICT = two resolvers name different metrics (must be 0).")
    write_sheet(wb, "Gates", ["gate", "findings", "result", "details"],
                [[name, n, "PASS" if n == 0 else "FAIL",
                  {"resolution conflicts": [c["phrase"] for c in g["conflicts"]], "registry drift": g["drift"],
                   "hard-cut coverage gaps": g["uncovered"], "sliced bindings": g["sliced"],
                   "live filters": g["filters"], "duplicate numbers": [f"{a} = {b}" for a, b, _ in g["duplicates"]]
                   }.get(name, [])] for name, n in g["gates"]], {1: 28, 4: 100})

    # ---- Lineage
    lrows = []
    used = sorted({o for v in serve_of.values() for o in v})
    for o in used:
        fq = f"serve.{o}"
        info = objs.get(fq, {})
        lrows.append([fq, info.get("engine", ""), info.get("refs", []), sorted(gold_sources(fq, objs)),
                      sorted(mid for mid, os_ in serve_of.items() if o in os_)])
    write_sheet(wb, "Lineage", ["serve_object", "engine", "direct_refs", "gold_sources", "metrics"], lrows,
                {1: 36, 3: 60, 4: 70, 5: 80},
                note="serve objects Cube v2 reads (from the generated SQL) and their upstream objects (ClickHouse DDL).")

    # ---- README
    sha = subprocess.run(["git", "-C", str(ROOT), "rev-parse", "--short", "HEAD"], capture_output=True,
                         text=True).stdout.strip()
    lines = [
        ("Seleric metric inventory — semantic v2 (live agent surface)", True),
        (f"Generated {datetime.datetime.now(datetime.timezone.utc):%Y-%m-%d %H:%M} UTC by "
         f"scripts/build_metric_inventory_v2.py (Seleric_Agent_Core {sha}).", False),
        (f"catalogue_v2 {svc.version}: {len(cat.metrics)} metrics, {len(cat.views)} views, {len(cat.dimensions)} "
         f"dimensions, {len(cat.hierarchies)} hierarchies, {len(cat.concepts)} concepts, {len(cat.glossary)} glossary "
         f"terms, {len(rrows)} retired v1 ids; agent registry {len(reg)} entries.", False),
        (f"Live values: Cube v2 {V2_URL}, brand {args.brand}, {s[:7]} and {p_s[:7]} (IST)." if live
         else "Live values skipped (--no-live).", False),
        ("Gates: " + "; ".join(f"{n} {'PASS' if k == 0 else f'FAIL ({k})'}" for n, k in g["gates"]), False),
        ("", False),
        ("Rules (doc/semantic_v2/PLAN.md §4): one number = one id = one home fact; scope / platform / paid are "
         "filters (Metric Variants); v1 ids are rejected naming their replacement (Retired v1 ids); bindings route "
         "hour → hourly fact and audience breakdowns → the Meta breakdown fact with ONE breakdown_type pinned.", False),
        ("Chain: gold → serve (ClickHouse DEFINER views / semantic.* snapshots) → Cube v2 (cube_serve login) → "
         "MCP catalogue_v2 (one resolver) → Seleric_Agent registry.", False),
        ("Replaces the v1 inventory (deprecated/doc/METRIC_INVENTORY_v1.xlsx).", False),
    ]
    from openpyxl.styles import Font
    for i, (text, bold) in enumerate(lines, 1):
        readme.cell(i, 1, text).font = Font(bold=bold, size=13 if bold else 11)
    readme.column_dimensions["A"].width = 160

    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"wrote {args.out}: {len(cat.metrics)} metrics, {len(vrows)} variants, {len(rrows)} retired ids, "
          f"{len(lrows)} serve objects; gates " + ", ".join(f"{n}={k}" for n, k in g["gates"]))
    return 1 if any(k for _, k in g["gates"]) else 0


if __name__ == "__main__":
    sys.exit(main())
